from pathlib import Path
from typing import List, Dict, Any
from pptx import Presentation
from .base import BaseDocumentProcessor, ProcessedDocument


class PPTXProcessor(BaseDocumentProcessor):
    """Processor for PowerPoint (PPTX) documents"""

    def __init__(self):
        super().__init__()
        self.supported_extensions = ['.pptx', '.ppt']

    def validate_file(self, file_path: Path) -> bool:
        """Validate if the PPTX file can be processed"""
        if not file_path.exists():
            return False

        if file_path.suffix.lower() not in ['.pptx']:
            # Note: python-pptx only supports .pptx, not .ppt
            return False

        try:
            Presentation(file_path)
            return True
        except Exception:
            return False

    def process(self, file_path: Path) -> ProcessedDocument:
        """
        Process a PowerPoint document and extract its content

        Args:
            file_path: Path to the PPTX file

        Returns:
            ProcessedDocument with extracted text and metadata
        """
        if not self.validate_file(file_path):
            raise ValueError(f"Invalid PPTX file: {file_path}")

        prs = Presentation(file_path)

        # Extract metadata
        metadata = self.extract_metadata(file_path)
        core_props = prs.core_properties

        metadata.update({
            "title": core_props.title or "",
            "author": core_props.author or "",
            "subject": core_props.subject or "",
            "keywords": core_props.keywords or "",
            "comments": core_props.comments or "",
            "num_slides": len(prs.slides),
        })

        # Extract text from all slides
        slides = []
        all_text = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_text = self._extract_slide_text(slide)
            slide_title = self._extract_slide_title(slide)

            slides.append({
                "slide_number": slide_num,
                "title": slide_title,
                "text": slide_text,
                "char_count": len(slide_text)
            })

            # Format slide text with title
            if slide_title:
                formatted_text = f"Slide {slide_num}: {slide_title}\n{slide_text}"
            else:
                formatted_text = f"Slide {slide_num}\n{slide_text}"

            all_text.append(formatted_text)

        combined_text = "\n\n".join(all_text)

        return ProcessedDocument(
            text=combined_text,
            metadata=metadata,
            pages=slides
        )

    def _extract_slide_text(self, slide) -> str:
        """Extract all text from a slide"""
        text_runs = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    text_runs.append(text)

        return "\n".join(text_runs)

    def _extract_slide_title(self, slide) -> str:
        """Extract the title from a slide"""
        if slide.shapes.title:
            return slide.shapes.title.text.strip()
        return ""

    def extract_slides_with_notes(self, file_path: Path) -> List[Dict[str, Any]]:
        """
        Extract slides with speaker notes

        Args:
            file_path: Path to the PPTX file

        Returns:
            List of dictionaries containing slide content and notes
        """
        if not self.validate_file(file_path):
            raise ValueError(f"Invalid PPTX file: {file_path}")

        prs = Presentation(file_path)
        slides_data = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_text = self._extract_slide_text(slide)
            slide_title = self._extract_slide_title(slide)

            # Extract speaker notes
            notes_text = ""
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    notes_text = notes_slide.notes_text_frame.text.strip()

            slides_data.append({
                "slide_number": slide_num,
                "title": slide_title,
                "content": slide_text,
                "notes": notes_text
            })

        return slides_data
